"""Licensed/manual JCC static data review. Uploads are drafts, never activated automatically."""

import csv
import hashlib
import io
import json
from datetime import datetime, timezone
from pathlib import Path
import zipfile

from openpyxl import load_workbook
from pydantic import Field, HttpUrl

from app.schemas.domain import Strict, Ruleset
from app.data_validation.catalog import validate_catalog

COLUMNS = [
    "entity_type",
    "id",
    "zh_CN",
    "en_US",
    "icon_url",
    "cost",
    "trait_ids",
    "thresholds",
    "trait_id",
    "verification_status",
    "evidence",
    "reviewer",
    "game",
    "patch",
    "set",
    "mode",
    "unsupported_reason",
]
MAX_BYTES = 2 * 1024 * 1024


class ManualMetadata(Strict):
    game: str = Field(pattern="^JCC$")
    ruleset_id: str = Field(pattern=r"^[a-zA-Z0-9_-]{1,100}$")
    patch: str = Field(min_length=1, max_length=100)
    set: str = Field(min_length=1, max_length=100)
    set_name: str = Field(min_length=1, max_length=200)
    mode: str = Field(min_length=1, max_length=100)
    source: str = Field(min_length=1, max_length=1000)
    source_url: HttpUrl
    license_status: str = Field(pattern="^(original|authorized|unknown)$")
    attribution: str = Field(min_length=1, max_length=1000)
    required_champion_ids: list[str] = Field(min_length=1, max_length=500)


def parse_rows(content, filename):
    if not content or len(content) > MAX_BYTES:
        raise ValueError("File must be nonempty and <=2 MiB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True))
    elif suffix == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 200
                or sum(e.file_size for e in entries) > 20 * 1024 * 1024
                or any("vbaproject" in e.filename.lower() or "externallinks/" in e.filename.lower() for e in entries)
            ):
                raise ValueError("Unsafe XLSX package")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.sheetnames) != 1:
                raise ValueError("Exactly one worksheet required")
            sheet = workbook.active
            if (sheet.max_row or 0) > 1001 or (sheet.max_column or 0) > len(COLUMNS):
                raise ValueError("File exceeds row/column limits")
            rows = []
            for cells in sheet.iter_rows():
                if len(rows) >= 1001 or len(cells) > len(COLUMNS) or any(c.data_type in ("f", "e") for c in cells):
                    raise ValueError("Formula/error/oversized worksheet")
                rows.append(["" if c.value is None else str(c.value) for c in cells])
        finally:
            workbook.close()
    else:
        raise ValueError("Only UTF-8 CSV / XLSX supported")
    if not rows or len(rows) > 1001 or rows[0] != COLUMNS:
        raise ValueError("Missing/invalid header or too many rows")
    result = []
    for values in rows[1:]:
        if not any(values):
            continue
        if len(values) != len(COLUMNS) or any(
            len(v) > 4000 or "\x00" in v or v.lstrip().startswith(("=", "@", "+")) for v in values
        ):
            raise ValueError("Invalid text row or spreadsheet formula")
        result.append(dict(zip(COLUMNS, (v.strip() for v in values))))
    if not result:
        raise ValueError("No records to validate")
    return result


def import_manual(content, filename, metadata):
    rows = parse_rows(content, filename)
    imported = datetime.now(timezone.utc).isoformat()
    base = {key: getattr(metadata, key) for key in ["game", "patch", "set", "mode", "source"]}
    base.update(imported_at=imported, verification_status="unverified")
    groups = {"champion": [], "trait": [], "emblem": []}
    reviewers = []
    for number, row in enumerate(rows, 2):
        if row["entity_type"] not in groups or not row["id"] or not row["zh_CN"] or not row["en_US"]:
            raise ValueError(f"Row {number}: invalid kind/ID/names")
        if any(row[k] != base[k] for k in ["game", "patch", "set", "mode"]):
            raise ValueError(f"Row {number}: mixed game/patch/set/mode")
        state = row["verification_status"].upper()
        if state not in ["VERIFIED", "PARTIALLY_VERIFIED", "UNVERIFIED", "UNSUPPORTED"]:
            raise ValueError(f"Row {number}: invalid real-data verification state")
        evidence = row["evidence"].split("|") if row["evidence"] else []
        if state == "VERIFIED" and (not evidence or not row["reviewer"] or metadata.license_status == "unknown"):
            raise ValueError(f"Row {number}: VERIFIED requires reviewer, evidence and known authorization")
        entity = dict(
            **base,
            id=row["id"],
            zh_CN=row["zh_CN"],
            en_US=row["en_US"],
            icon_path=row["icon_url"],
            verification_evidence=evidence,
        )
        entity["verification_status"] = state.lower()
        verified = state == "VERIFIED"
        if row["entity_type"] == "champion":
            entity.update(
                cost=int(row["cost"]),
                traits=row["trait_ids"].split("|") if row["trait_ids"] else [],
                available=verified and not row["unsupported_reason"],
                unlock_condition=row["unsupported_reason"] or None,
            )
        elif row["entity_type"] == "trait":
            counts = [int(v) for v in row["thresholds"].split("|")] if row["thresholds"] else []
            entity.update(
                thresholds=[dict(count=n, utility=i + 1) for i, n in enumerate(counts)],
                supported=verified and not row["unsupported_reason"],
                special_conditions=row["unsupported_reason"] or (None if verified else "Unverified manual mechanic"),
            )
        else:
            entity.update(trait_id=row["trait_id"], available=verified and not row["unsupported_reason"])
        groups[row["entity_type"]].append(entity)
        reviewers.append(dict(row=number, id=row["id"], reviewer=row["reviewer"], evidence=evidence))
    fully_reviewed = all(r["verification_status"].upper() == "VERIFIED" for r in rows)
    readiness = dict(
        source_kind="manual_review",
        set_name=metadata.set_name,
        resource_patch=metadata.patch,
        build_version="manual",
        last_verified_at=imported if fully_reviewed else None,
        live_as_of=None,
        optimization_enabled=False,
        scope="人工校验草稿；完整游戏规则未经独立审查，禁止实战优化",
        completeness=sum(r["verification_status"].upper() == "VERIFIED" for r in rows) / len(rows),
        required_champion_ids=metadata.required_champion_ids,
        snapshots=[
            dict(
                filename=Path(filename).name,
                url=str(metadata.source_url),
                sha256=hashlib.sha256(content).hexdigest(),
                imported_at=imported,
                attribution=metadata.attribution,
                reviewers=reviewers,
            )
        ],
        rule_checks=[
            dict(
                id="jcc_game_rules",
                status="UNVERIFIED",
                essential=True,
                evidence=[],
                notes="No authorized complete ruleset configured",
            )
        ],
        unsupported_mechanics=["金铲铲独立版本、特殊规则及完整名单尚未审核；禁止套用 TFT 规则"],
    )
    raw = dict(
        **base,
        id=metadata.ruleset_id,
        zh_CN=metadata.set_name,
        en_US=metadata.set_name,
        champions=groups["champion"],
        traits=groups["trait"],
        emblems=groups["emblem"],
        allow_duplicate_emblems=False,
        allow_unused_emblems=True,
        license_status=metadata.license_status,
        data_readiness=readiness,
    )
    raw["verification_status"] = "partially_verified" if fully_reviewed else "unverified"
    catalog = Ruleset.model_validate(raw)
    report = validate_catalog(catalog.model_dump())
    return dict(catalog=catalog.model_dump(), validation=report, review_required=True, activated=False)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("file")
    parser.add_argument("--metadata", required=True)
    parser.add_argument("--output", required=True, help="Review draft JSON, never installed automatically")
    args = parser.parse_args()
    metadata = ManualMetadata.model_validate_json(Path(args.metadata).read_text())
    result = import_manual(Path(args.file).read_bytes(), args.file, metadata)
    Path(args.output).write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result["validation"], ensure_ascii=False, indent=2))
