"""Bounded, offline community guide import. Never consumed by the optimizer."""

import csv
import hashlib
import io
import json
import unicodedata
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, HttpUrl
from sqlalchemy import select, update
from sqlalchemy.exc import IntegrityError

from app.models.database import Session
from app.models.community import GuideImport, GuideRow
from app.repositories.catalog import get_catalog
from app.rules.evaluator import FixtureRules


class Attribution(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    game_id: str
    ruleset_id: str
    source_title: str = Field(min_length=1, max_length=300)
    source_url: HttpUrl | None = None
    attribution: str = Field(min_length=1, max_length=1000)
    license_status: str = Field(pattern="^(original|authorized|unknown)$")
    sheet: str | None = Field(default=None, max_length=100)


class Correction(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    expected_revision: int = Field(ge=1)
    emblem_name: str = Field(max_length=1000)
    emblem_icon_ref: str = Field(default="", max_length=1000)
    recommended_holders: str = Field(default="", max_length=1000)
    composition_aliases: str = Field(default="", max_length=1000)
    notes: str = Field(default="", max_length=1000)
    reviewer: str = Field(min_length=1, max_length=200)
    correction_note: str = Field(min_length=1, max_length=1000)
    confirm: bool = False


COLUMNS = {"emblem_name", "emblem_icon_ref", "recommended_holders", "composition_aliases", "label_status", "notes"}
MAX_BYTES = 5 * 1024 * 1024


def now():
    return datetime.now(timezone.utc).isoformat()


def parse_file(content, filename, sheet=None):
    if not content or len(content) > MAX_BYTES:
        raise ValueError("文件必须非空且不超过 5 MiB")
    suffix = Path(filename).suffix.lower()
    selected = None
    try:
        if suffix == ".csv":
            if sheet:
                raise ValueError("CSV 不支持工作表选择")
            rows = csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True)
        elif suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 200 or sum(e.file_size for e in entries) > 20 * 1024 * 1024:
                    raise ValueError("XLSX 解压大小超出限制")
                if any("externallinks/" in e.filename.lower() or "vbaproject" in e.filename.lower() for e in entries):
                    raise ValueError("不支持外部链接或宏")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if sheet is None and len(workbook.sheetnames) != 1:
                    raise ValueError("多工作表文件必须指定 sheet")
                selected = sheet or workbook.sheetnames[0]
                if selected not in workbook.sheetnames:
                    raise ValueError("工作表不存在")
                ws = workbook[selected]
                if (ws.max_row or 0) > 1001 or (ws.max_column or 0) > len(COLUMNS):
                    raise ValueError("最多 1000 行数据、6 列")
                values = []
                for number, cells in enumerate(ws.iter_rows(), 1):
                    if number > 1001 or len(cells) > len(COLUMNS):
                        raise ValueError("最多 1000 行数据、6 列")
                    if any(c.data_type in ("f", "e") for c in cells):
                        raise ValueError(f"第 {number} 行不支持公式或错误单元格")
                    values.append([c.value if c.value is not None else "" for c in cells])
                rows = iter(values)
            finally:
                workbook.close()
        else:
            raise ValueError("仅支持 UTF-8 CSV 或 XLSX")
        header = next(rows, None)
        if not header or any(not isinstance(c, str) for c in header):
            raise ValueError("缺少文本表头")
        header = [c.strip() for c in header]
        if len(set(header)) != len(header) or not set(header) <= COLUMNS or "emblem_name" not in header:
            raise ValueError("表头重复、未知或缺少 emblem_name；统计字段不受支持")
        result = []
        for number, values in enumerate(rows, 2):
            if number > 1001:
                raise ValueError("最多 1000 行数据")
            if not any(v != "" for v in values):
                continue
            if len(values) != len(header):
                raise ValueError(f"第 {number} 行列数不匹配")
            if any(
                not isinstance(v, str) or len(v) > 1000 or "\x00" in v or v.lstrip().startswith("=") for v in values
            ):
                raise ValueError(f"第 {number} 行必须为长度不超过 1000 的文本，禁止公式")
            row = dict(zip(header, values))
            if not row.get("emblem_name", "").strip() and not row.get("emblem_icon_ref", "").strip():
                raise ValueError(f"第 {number} 行缺少纹章名称或图标引用")
            if row.get("label_status", "").strip() not in ("", "unverified", "uncertain"):
                raise ValueError(f"第 {number} 行 label_status 无效")
            result.append((number, row))
        if not result:
            raise ValueError("没有可导入的数据行")
        return result, selected
    except (UnicodeError, csv.Error, zipfile.BadZipFile, KeyError) as exc:
        raise ValueError(f"文件无法解析: {exc}") from exc
    except ValueError:
        raise
    except Exception as exc:
        # Parser failures are invalid uploads; never persist a partially parsed file.
        raise ValueError("XLSX 文件无法解析") from exc


def normalize(value):
    return unicodedata.normalize("NFKC", value).strip().casefold()


def tokens(value):
    return list(dict.fromkeys(v.strip() for v in value.split("|") if v.strip()))


def resolve(label, entities):
    matches = [e for e in entities if normalize(label) in {normalize(v) for v in [e.id, e.zh_CN, e.en_US, *e.aliases]}]
    return matches[0] if len(matches) == 1 else None


def evaluate(row, catalog, confirmed=False):
    name = row.get("emblem_name", "")
    emblem = resolve(name, catalog.emblems) if name else None
    issues = []
    if not emblem:
        issues.append("纹章名称未唯一匹配；图标引用不会自动识别")
    holders = []
    for label in tokens(row.get("recommended_holders", "")):
        holder = resolve(label, catalog.champions)
        if holder is None:
            issues.append(f"推荐持有者未唯一匹配: {label}")
        elif emblem and not FixtureRules().legal(holder, emblem):
            issues.append(f"推荐持有者不符合当前规则: {label}")
        holders.append({"label": label, "champion_id": holder.id if holder else None})
    if row.get("label_status", "").strip() == "uncertain" and not confirmed:
        issues.append("来源标记为不确定，等待人工确认")
    return {
        **row,
        "emblem_id": emblem.id if emblem else None,
        "holders": holders,
        "aliases": tokens(row.get("composition_aliases", "")),
        "issues": issues,
        "status": "needs_review" if issues else ("reviewed" if confirmed else "matched"),
    }


def serialize(row):
    return dict(
        id=row.id,
        row_number=row.row_number,
        revision=row.revision,
        raw=row.raw,
        current=row.current,
        history=row.history,
    )


def get_import(import_id):
    with Session() as session:
        batch = session.get(GuideImport, import_id)
        if batch is None:
            raise LookupError("导入记录不存在")
        rows = session.scalars(
            select(GuideRow).where(GuideRow.import_id == import_id).order_by(GuideRow.row_number)
        ).all()
        return dict(id=batch.id, **batch.metadata_json, rows=[serialize(r) for r in rows])


def list_imports(game_id, ruleset_id):
    get_catalog(game_id, ruleset_id)
    with Session() as session:
        batches = session.scalars(select(GuideImport)).all()
        return [
            dict(id=b.id, **b.metadata_json)
            for b in batches
            if b.metadata_json["game_id"] == game_id and b.metadata_json["ruleset_id"] == ruleset_id
        ]


def import_guide(content, filename, attribution):
    catalog = get_catalog(attribution.game_id, attribution.ruleset_id)
    parsed, sheet = parse_file(content, filename, attribution.sheet)
    metadata = attribution.model_dump(mode="json")
    metadata.update(
        catalog_sha256=hashlib.sha256(catalog.model_dump_json().encode()).hexdigest(),
        sheet=sheet,
        filename=Path(filename).name,
        file_sha256=hashlib.sha256(content).hexdigest(),
        patch=catalog.patch,
        set=catalog.set,
        mode=catalog.mode,
        kind="community_recommendation",
        verification_status="unverified",
    )
    identity = hashlib.sha256(json.dumps(metadata, sort_keys=True).encode()).hexdigest()
    try:
        with Session.begin() as session:
            if session.get(GuideImport, identity) is None:
                metadata.update(imported_at=now(), row_count=len(parsed))
                session.add(GuideImport(id=identity, metadata_json=metadata))
                session.flush()
                session.add_all(
                    [
                        GuideRow(
                            id=f"{identity}:{number}",
                            import_id=identity,
                            row_number=number,
                            raw=row,
                            current=evaluate(row, catalog),
                            history=[],
                        )
                        for number, row in parsed
                    ]
                )
    except IntegrityError:
        # Concurrent identical imports may race on the unique batch key.
        with Session() as session:
            if session.get(GuideImport, identity) is None:
                raise
    return get_import(identity)


def correct_row(import_id, row_id, correction):
    with Session.begin() as session:
        batch = session.get(GuideImport, import_id)
        row = session.get(GuideRow, row_id)
        if batch is None or row is None or row.import_id != import_id:
            raise LookupError("导入记录或行不存在")
        if row.revision != correction.expected_revision:
            raise RuntimeError("记录已被修改，请重新加载")
        catalog = get_catalog(batch.metadata_json["game_id"], batch.metadata_json["ruleset_id"])
        if hashlib.sha256(catalog.model_dump_json().encode()).hexdigest() != batch.metadata_json.get("catalog_sha256"):
            raise RuntimeError("目录内容已改变，请使用新的规则版本重新导入")
        fields = correction.model_dump(exclude={"expected_revision", "reviewer", "correction_note", "confirm"})
        if not fields["emblem_name"] and not fields["emblem_icon_ref"]:
            raise ValueError("缺少纹章名称或图标引用")
        fields["label_status"] = row.raw.get("label_status", "unverified")
        current = evaluate(fields, catalog, correction.confirm)
        if correction.confirm and current["issues"]:
            raise ValueError("仍存在未解决的标签或非法持有者，无法确认")
        history = [
            *row.history,
            dict(
                at=now(),
                reviewer=correction.reviewer,
                note=correction.correction_note,
                before=row.current,
                after=current,
                revision=row.revision + 1,
            ),
        ]
        changed = session.execute(
            update(GuideRow)
            .where(GuideRow.id == row_id, GuideRow.revision == correction.expected_revision)
            .values(current=current, history=history, revision=row.revision + 1)
        )
        if changed.rowcount != 1:
            raise RuntimeError("记录已被修改，请重新加载")
    return get_import(import_id)


def main():
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("--game", required=True, choices=["TFT", "JCC"])
    parser.add_argument("--ruleset", required=True)
    parser.add_argument("--source-title", required=True)
    parser.add_argument("--attribution", required=True)
    parser.add_argument("--license-status", required=True, choices=["original", "authorized", "unknown"])
    parser.add_argument("--source-url")
    parser.add_argument("--sheet")
    args = parser.parse_args()
    try:
        with args.file.open("rb") as file:
            content = file.read(MAX_BYTES + 1)
        batch = import_guide(
            content,
            args.file.name,
            Attribution(
                game_id=args.game,
                ruleset_id=args.ruleset,
                source_title=args.source_title,
                attribution=args.attribution,
                license_status=args.license_status,
                source_url=args.source_url,
                sheet=args.sheet,
            ),
        )
    except (ValueError, OSError) as exc:
        parser.error(str(exc))
    print(json.dumps(batch, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
